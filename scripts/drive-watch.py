#!/usr/bin/env python3
"""Drive folder watcher: list recursively, diff vs manifest, download new files
(after a grace period so in-progress uploads finish), quick-profile them.

Stdout: Vietnamese summary of newly downloaded files, or the line NO_NEW_FILES.
"""
import json, os, re, subprocess, sys, time, urllib.request

ROOT_ID = "1Cf43QiMnao0ubHHLtcfg5uBHZhbD2Svw"
BASE = os.environ.get("DRIVE_WATCH_DIR", os.path.dirname(os.path.abspath(__file__)))
FILES_DIR = os.path.join(BASE, "files")
MANIFEST = os.path.join(BASE, "manifest.json")
VENV_PY = os.environ.get("DW_VENV_PY", os.path.join(BASE, "venv", "bin", "python"))
GRACE_SECS = 15 * 60  # do not download files first seen less than 15 min ago


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=40) as r:
        return r.read().decode("utf-8", errors="replace")


def list_folder(fid):
    html = fetch(f"https://drive.google.com/embeddedfolderview?id={fid}#list")
    entries = []
    for p in html.split('<div class="flip-entry"')[1:]:
        m = re.search(r'id="entry-([A-Za-z0-9_-]+)"', p)
        if not m:
            continue
        t = re.search(r'class="flip-entry-title">(.*?)</div>', p)
        mod = re.search(r'flip-entry-last-modified"><div>(.*?)</div>', p)
        entries.append({
            "id": m.group(1),
            "name": (t.group(1).strip() if t else m.group(1)),
            "is_folder": 'aria-label="Folder"' in p,
            "modified": (mod.group(1).strip() if mod else ""),
        })
    return entries


def walk(fid, prefix=""):
    out = []
    for e in list_folder(fid):
        path = f"{prefix}/{e['name']}" if prefix else e["name"]
        rec = dict(e, path=path)
        out.append(rec)
        if e["is_folder"]:
            out.extend(walk(e["id"], path))
    return out


def safe_path(path):
    # keep hierarchy but sanitize each component
    parts = [re.sub(r'[^\w\-. \(\)\[\]]+', '_', c, flags=re.U).strip(" .") or "_"
             for c in path.split("/")]
    return os.path.join(*parts)


def download_file(fid, dest):
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    url = f"https://drive.google.com/uc?id={fid}"
    r = subprocess.run([VENV_PY, "-m", "gdown", "-O", dest, url],
                       capture_output=True, text=True, timeout=1800)
    return r.returncode == 0 and os.path.exists(dest) and os.path.getsize(dest) > 0


def profile_excel(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            try:
                rows = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                headers = [str(v)[:40] for v in (rows[0] if rows else []) if v is not None][:12]
            except Exception:
                headers = []
            sheets.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                           "headers": headers})
        wb.close()
        return {"type": "excel", "sheets": sheets}
    except Exception as e:
        return {"type": "excel", "error": str(e)[:200]}


def profile_file(path):
    ext = os.path.splitext(path)[1].lower()
    size_mb = os.path.getsize(path) / 1048576
    info = {"size_mb": round(size_mb, 1)}
    try:
        if ext in (".xlsx", ".xlsm"):
            info.update(profile_excel(path))
        elif ext == ".xls":
            info["type"] = "excel-legacy"
        elif ext == ".csv":
            import csv
            with open(path, encoding="utf-8-sig", errors="replace") as f:
                rows = list(csv.reader(f))
            info.update({"type": "csv", "rows": len(rows),
                         "headers": [str(v)[:40] for v in rows[0]][:12] if rows else []})
        elif ext == ".pdf":
            info["type"] = "pdf"
        elif ext in (".png", ".jpg", ".jpeg", ".webp"):
            from PIL import Image
            im = Image.open(path)
            info.update({"type": "image", "dimensions": f"{im.width}x{im.height}"})
        else:
            info["type"] = ext.lstrip(".") or "unknown"
    except Exception as e:
        info["error"] = str(e)[:200]
    return info


def main():
    force = "--force" in sys.argv  # bypass grace period (user confirmed upload done)
    os.makedirs(FILES_DIR, exist_ok=True)
    manifest = json.load(open(MANIFEST, encoding="utf-8")) if os.path.exists(MANIFEST) else {}
    now = time.time()
    try:
        entries = walk(ROOT_ID)
    except Exception as e:
        print(f"LIST_ERROR: {e}")
        return 1
    files = [e for e in entries if not e["is_folder"]]

    for e in files:
        rec = manifest.get(e["id"])
        if not rec:
            manifest[e["id"]] = {"name": e["name"], "path": e["path"],
                                 "first_seen": now, "downloaded": False, "failed": 0}

    newly = []
    for e in files:
        rec = manifest[e["id"]]
        if rec.get("downloaded") or rec.get("failed", 0) >= 3:
            continue
        if now - rec["first_seen"] < GRACE_SECS and not force:
            continue  # possibly still uploading
        dest = os.path.join(FILES_DIR, safe_path(e["path"]))
        ok = download_file(e["id"], dest)
        if ok:
            rec["downloaded"] = True
            rec["local"] = os.path.relpath(dest, BASE)
            prof = profile_file(dest)
            newly.append({"name": e["name"], "path": e["path"], "profile": prof})
        else:
            rec["failed"] = rec.get("failed", 0) + 1

    json.dump(manifest, open(MANIFEST, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    if not newly:
        print("NO_NEW_FILES")
        return 0

    lines = [f"PHAT HIEN {len(newly)} FILE MOI TREN DRIVE (PFC2026):", ""]
    for n in newly:
        p = n["profile"]
        lines.append(f"- {n['path']} ({p.get('size_mb', '?')} MB)")
        if p.get("type") == "excel":
            if p.get("error"):
                lines.append(f"    [excel] loi doc: {p['error']}")
            else:
                for s in p["sheets"]:
                    hdr = " | ".join(s["headers"]) if s["headers"] else "(khong doc duoc header)"
                    lines.append(f"    [sheet] {s['name']}: {s['rows']} dong x {s['cols']} cot")
                    lines.append(f"      cot: {hdr}")
        elif p.get("type") == "csv":
            hdr = " | ".join(p["headers"]) if p.get("headers") else ""
            lines.append(f"    [csv] {p['rows']} dong. cot: {hdr}")
        elif p.get("type"):
            lines.append(f"    [{p['type']}]" + (f" {p['dimensions']}" if p.get("dimensions") else ""))
        if p.get("error") and p.get("type") != "excel":
            lines.append(f"    loi profile: {p['error']}")
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
