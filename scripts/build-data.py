#!/usr/bin/env python3
"""Build static data for the PFC Visual Browser web app.

Scans every .xlsx under drive-watch/files/, reads each sheet's structure
(openpyxl read_only), extracts embedded images from xl/media, dedupes by
full md5, resizes to max 900px and saves as JPEG q70 under public/img/.

Outputs:
  <web>/public/data/index.json   - brands/files/sheets structure
  <web>/public/data/images.json   - md5 -> {w, h, files[]}
  <web>/public/img/<md5>.jpg      - deduped resized images

Run: /home/hatch/workspace/drive-watch/venv/bin/python scripts/build-data.py
"""
import hashlib
import io
import json
import os
import re
import sys
import zipfile

WEB_DIR = os.environ.get("WEB_DIR", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC_DIR = os.environ.get("DRIVE_FILES_DIR", "/home/hatch/workspace/drive-watch/files")
OUT_DATA = os.path.join(WEB_DIR, "public", "data")
OUT_IMG = os.path.join(WEB_DIR, "public", "img")
MAX_DIM = 900
JPEG_Q = 70
HEADER_COLS = 15


def slugify(name: str) -> str:
    s = name.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    s = re.sub(r"-{2,}", "-", s)
    return s or "file"


def hinh_the_for(group: str | None, filename: str) -> str:
    """Hinh the (product silhouette): subfolder name, or parsed from filename.

    e.g. JILEON/ANKLE BOOTS/*.xlsx -> "ANKLE BOOTS";
    260627_PFC_ADIDAS RAINBOOT W.xlsx -> "RAINBOOT W".
    """
    if group:
        return group
    m = re.match(r"^\d+_PFC_[A-Za-z]+\s+(.+)\.xlsx?$", filename, re.I)
    if m:
        return re.sub(r"\s+", " ", m.group(1)).strip().upper()
    return "KHÁC"


def read_sheets(path: str):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            try:
                rows = ws.max_row or 0
                cols = ws.max_column or 0
                headers = []
                try:
                    # first row with >=2 non-empty cells within first 8 rows
                    for row in ws.iter_rows(min_row=1, max_row=8,
                                           values_only=True):
                        vals = ["" if v is None else str(v)[:80]
                                for v in list(row)[:HEADER_COLS]]
                        if sum(1 for v in vals if v.strip()) >= 2:
                            headers = vals
                            break
                except Exception:
                    pass
                sheets.append({"name": ws.title, "rows": rows,
                               "cols": cols, "headers": headers})
            except Exception as e:  # noqa: BLE001
                sheets.append({"name": ws.title, "rows": 0, "cols": 0,
                               "headers": [], "error": str(e)[:120]})
    finally:
        wb.close()
    return sheets


def extract_media(path: str):
    """Return list of (md5, bytes) for readable raster images in xl/media."""
    out = []
    try:
        z = zipfile.ZipFile(path)
    except Exception:
        return out
    try:
        names = [n for n in z.namelist() if n.startswith("xl/media/")]
        for n in names:
            ext = os.path.splitext(n)[1].lower()
            if ext not in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff"):
                continue  # skip .emf/.wmf vector formats
            try:
                data = z.read(n)
            except Exception:
                continue
            if len(data) < 100:
                continue
            out.append((hashlib.md5(data).hexdigest(), data))
    finally:
        z.close()
    return out


def save_image(md5: str, data: bytes):
    """Resize to MAX_DIM, save JPEG q70. Returns (w, h) original or None."""
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception:
        return None
    w, h = im.size
    if im.mode in ("RGBA", "LA", "PA"):
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[-1])
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")
    if max(w, h) > MAX_DIM:
        scale = MAX_DIM / max(w, h)
        im = im.resize((int(w * scale), int(h * scale)), Image.LANCZOS)
    dest = os.path.join(OUT_IMG, f"{md5}.jpg")
    im.save(dest, "JPEG", quality=JPEG_Q, optimize=True)
    return (w, h)


def main() -> int:
    os.makedirs(OUT_DATA, exist_ok=True)
    os.makedirs(OUT_IMG, exist_ok=True)

    brands = [
        {"id": "adidas", "name": "ADIDAS", "dir": os.path.join(SRC_DIR, "ADIDAS")},
        {"id": "jileon", "name": "JILEON", "dir": os.path.join(SRC_DIR, "JILEON")},
    ]

    image_store = {}   # md5 -> {"w":, "h":, "files": set}
    image_bytes = {}   # md5 -> raw bytes (first seen)
    out_brands = []
    total_sheets = 0

    for b in brands:
        files_out = []
        xlsx_paths = []
        for root, _d, fs in os.walk(b["dir"]):
            for f in sorted(fs):
                if f.lower().endswith((".xlsx", ".xlsm")):
                    xlsx_paths.append(os.path.join(root, f))
        for path in xlsx_paths:
            rel = os.path.relpath(path, b["dir"])
            stem = os.path.splitext(os.path.basename(path))[0]
            group = os.path.dirname(rel) or None
            fid = slugify(stem)
            size_mb = round(os.path.getsize(path) / 1048576, 1)
            print(f"[{b['id']}] {rel} ({size_mb} MB)...", flush=True)
            sheets = read_sheets(path)
            total_sheets += len(sheets)
            media = extract_media(path)
            md5s = []
            for md5, data in media:
                md5s.append(md5)
                if md5 not in image_bytes:
                    image_bytes[md5] = data
                rec = image_store.setdefault(
                    md5, {"w": 0, "h": 0, "files": set()})
                rec["files"].add(f"{b['id']}/{fid}")
            files_out.append({
                "id": fid,
                "name": os.path.basename(path),
                "group": group,
                "hinh_the": hinh_the_for(group, os.path.basename(path)),
                "sizeMB": size_mb,
                "sheets": sheets,
                "images": sorted(set(md5s)),
            })
            print(f"    sheets={len(sheets)} images={len(set(md5s))}", flush=True)
        out_brands.append({"id": b["id"], "name": b["name"], "files": files_out})

    # dedupe + resize + save images
    print(f"Saving {len(image_bytes)} unique images...", flush=True)
    skipped = 0
    for i, (md5, data) in enumerate(image_bytes.items()):
        wh = save_image(md5, data)
        if wh is None:
            skipped += 1
            image_store.pop(md5, None)
            continue
        image_store[md5]["w"], image_store[md5]["h"] = wh
        if (i + 1) % 100 == 0:
            print(f"  {i + 1}/{len(image_bytes)}", flush=True)

    # prune image lists in files to successfully saved images
    ok = set(image_store)
    for b in out_brands:
        for f in b["files"]:
            f["images"] = [m for m in f["images"] if m in ok]

    images_json = {m: {"w": r["w"], "h": r["h"], "files": sorted(r["files"])}
                   for m, r in image_store.items()}
    json.dump({"brands": out_brands},
              open(os.path.join(OUT_DATA, "index.json"), "w", encoding="utf-8"),
              ensure_ascii=False)
    json.dump(images_json,
              open(os.path.join(OUT_DATA, "images.json"), "w", encoding="utf-8"),
              ensure_ascii=False)

    # totals
    n_files = sum(len(b["files"]) for b in out_brands)
    pub_size = 0
    for root, _d, fs in os.walk(os.path.join(WEB_DIR, "public")):
        for f in fs:
            if f == ".gitkeep":
                continue
            pub_size += os.path.getsize(os.path.join(root, f))
    print("=" * 50)
    print(f"brands=2 files={n_files} sheets={total_sheets} "
          f"unique_images={len(images_json)} skipped={skipped}")
    print(f"public/ size = {pub_size / 1048576:.1f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
