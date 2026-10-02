#!/usr/bin/env python3
"""Extract cong-doan (production stages) from PFC Excel files.

Patterns:
  LUU TRINH : block-based sheets (ADIDAS) - "LUU TRINH X" title,
              block 1 = "Ten chi tiet" detail table,
              blocks 2+ = cong doan (ten, hinh anh, thong so, chu y, dung cu, bao ho)
  DMTG      : "BANG DINH MUC THOI GIAN" tables (both brands) -
              STT | CONG DOAN | T GIAN | NGUOI | MAY MOC | GHI CHU
Images are mapped to cong-doan blocks via drawingML anchors (fast zipfile parse).
Sheet cells are batch-read once via iter_rows (read_only random access is slow).

Output: <web>/public/data/stages.json
"""
import os, re, json, hashlib, zipfile
from collections import Counter
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

SRC = "/home/hatch/workspace/drive-watch/files"
WEB = "/home/hatch/workspace/pfc-web"
IMG_INDEX = json.load(open(os.path.join(WEB, "public/data/images.json")))
OUT = os.path.join(WEB, "public/data/stages.json")

BRANDS = [("adidas", "ADIDAS"), ("jileon", "JILEON")]
STAGE_ORDER = ["CHẶT", "CÁN", "CÁN LUYỆN", "IN", "MAY", "GÒ", "THÀNH HÌNH",
               "HOÀN TẤT", "KCS"]

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def norm(v):
    return re.sub(r"\s+", " ", str(v).strip() if v is not None else "")


def slugify(name):
    import unicodedata
    s = unicodedata.normalize("NFD", name.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return re.sub(r"-{2,}", "-", s) or "x"


def stage_of(sheet_name):
    s = norm(sheet_name).upper()
    if "CÁN LUYỆN" in s:
        return "CÁN LUYỆN"
    if "CÁN-" in s or s.strip() == "CÁN":
        return "CÁN"
    if "CHẶT" in s:
        return "CHẶT"
    if s == "IN" or s.startswith("IN_") or s.startswith("QTCN IN"):
        return "IN"
    if "THÀNH HÌNH" in s:
        return "THÀNH HÌNH"
    if s == "MAY" or "QTCN MAY" in s:
        return "MAY"
    if "GÒ" in s:
        return "GÒ"
    if "HOÀN TẤT" in s:
        return "HOÀN TẤT"
    if "KCS" in s:
        return "KCS"
    return None


# ---------------------------------------------------------------- images ---
def sheet_image_anchors(path):
    """{sheet_title: [(md5, row0, col0)]} via drawingML, fast."""
    out = {}
    try:
        z = zipfile.ZipFile(path)
    except Exception:
        return out
    try:
        names = set(z.namelist())
        media_md5 = {}
        for n in names:
            if n.startswith("xl/media/"):
                try:
                    media_md5[n] = hashlib.md5(z.read(n)).hexdigest()
                except Exception:
                    pass
        try:
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            wb_rel = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        except Exception:
            return out
        rel_target = {}
        for rel in wb_rel.findall("pkgrel:Relationship", NS):
            rid = rel.get("Id")
            tgt = rel.get("Target", "")
            if tgt.startswith("worksheets/"):
                rel_target[rid] = "xl/" + tgt
        sheet_file = {}
        for sh in wb_xml.findall("main:sheets/main:sheet", NS):
            title = sh.get("name")
            rid = sh.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if title and rid in rel_target:
                sheet_file[title] = rel_target[rid]
        for title, sfile in sheet_file.items():
            try:
                srel = ET.fromstring(z.read(
                    "xl/worksheets/_rels/" + sfile.split("/")[-1] + ".rels"))
            except Exception:
                continue
            draw_file = None
            for rel in srel.findall("pkgrel:Relationship", NS):
                if rel.get("Type", "").endswith("/drawing"):
                    draw_file = rel.get("Target", "").replace("../", "xl/")
                    break
            if not draw_file or draw_file not in names:
                continue
            try:
                drel = ET.fromstring(z.read(
                    "xl/drawings/_rels/" + draw_file.split("/")[-1] + ".rels"))
            except Exception:
                continue
            embed_map = {}
            for rel in drel.findall("pkgrel:Relationship", NS):
                tgt = rel.get("Target", "").replace("../", "xl/")
                if tgt in media_md5:
                    embed_map[rel.get("Id")] = media_md5[tgt]
            try:
                draw = ET.fromstring(z.read(draw_file))
            except Exception:
                continue
            imgs = []
            for tag in ("xdr:twoCellAnchor", "xdr:oneCellAnchor"):
                for anc in draw.findall(tag, NS):
                    frm = anc.find("xdr:from", NS)
                    blip = anc.find("xdr:pic/xdr:blipFill/a:blip", NS)
                    if blip is None:  # shape without picture
                        continue
                    if frm is None:
                        continue
                    try:
                        r0 = int(frm.find("xdr:row", NS).text)
                        c0 = int(frm.find("xdr:col", NS).text)
                    except Exception:
                        continue
                    eid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    md5 = embed_map.get(eid)
                    if md5 and md5 in IMG_INDEX:
                        imgs.append((md5, r0, c0))
            if imgs:
                out[title] = imgs
    finally:
        z.close()
    return out


# ------------------------------------------------------------- batch read --
def read_all(ws, max_row, max_col):
    """1-indexed rows[r][c]; None for empty. Single sequential pass."""
    rows = [[None] * (max_col + 1) for _ in range(max_row + 1)]
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r > max_row:
            break
        rr = rows[r]
        for c, v in enumerate(row[:max_col], start=1):
            rr[c] = v
    return rows


def val(rows, r, c):
    try:
        return rows[r][c]
    except IndexError:
        return None


def nval(rows, r, c):
    return norm(val(rows, r, c))


# ------------------------------------------------------------ detectors ----
def detect_luutrinh(rows):
    for r in range(1, min(len(rows), 5)):
        for c in range(1, min(len(rows[r]), 40)):
            if nval(rows, r, c).upper().startswith("LƯU TRÌNH"):
                return True
    return False


def detect_dmtg(rows):
    for r in range(1, min(len(rows), 10)):
        for c in range(1, min(len(rows[r]), 6)):
            if "ĐỊNH MỨC THỜI GIAN" in nval(rows, r, c).upper():
                return True
    return False


# ------------------------------------------------------- luu trinh parse ---
def _int_or_none(s):
    try:
        return int(float(norm(s).replace(",", ".")))
    except Exception:
        return None


def parse_congdoan_block(rows, start, end, max_col, stt, ten):
    notes, specs, tools, ppe = [], [], [], []
    note_row = tool_row = ppe_row = spec_row = None
    spec_labels = []
    for r in range(start, end + 1):
        c31 = nval(rows, r, 31)
        c1 = nval(rows, r, 1)
        low31 = c31.lower()
        if note_row is None and low31.startswith("chú ý"):
            note_row = r
        if tool_row is None and "dụng cụ tools" in low31:
            tool_row = r
        if ppe_row is None and "bảo hộ" in low31:
            ppe_row = r
        if spec_row is None and r > start + 2:
            n = sum(1 for c in range(1, 31) if nval(rows, r, c))
            if n >= 4 and not c1.lower().startswith("hình ảnh"):
                spec_row = r
                spec_labels = [nval(rows, r, c) for c in range(1, 31)
                               if nval(rows, r, c)]
    nend = tool_row or ppe_row or (end + 1)

    def _sig(r):
        return "现场" in nval(rows, r, 1)

    if note_row:
        for r in range(note_row + 1, nend):
            if _sig(r):
                break
            v = nval(rows, r, 31)
            if v:
                notes.append(v)
    if spec_row and spec_labels:
        # pair label<->value by COLUMN (merged cells share the same columns)
        label_cols = [(c, nval(rows, spec_row, c)) for c in range(1, 31)
                      if nval(rows, spec_row, c)]
        for r in range(spec_row + 1, min(spec_row + 6, nend)):
            vals = [nval(rows, r, c) for c in range(1, 31)]
            if sum(1 for v in vals if v) < 2:
                break
            row = {}
            for c, lab in label_cols:
                v = nval(rows, r, c)
                if v:
                    row[lab] = v
            if row:
                specs.append(row)
    if tool_row:
        tend = ppe_row or (end + 1)
        for r in range(tool_row + 1, tend):
            v = nval(rows, r, 31)
            if v and not nval(rows, r, 1):
                tools.append(v)
    if ppe_row:
        for r in range(ppe_row + 1, end + 1):
            if nval(rows, r, 1):
                break
            v = nval(rows, r, 31)
            if v:
                ppe.append(v)
    return {"stt": stt, "ten": ten, "notes": notes, "specs": specs,
            "tools": tools, "ppe": ppe, "_start": start, "_end": end}


def _is_reference_block(rows, start, end):
    """Block 1 style: a detail/CCDC table, not a cong doan (no real name)."""
    for rr in range(start, min(end, start + 30) + 1):
        up = " ".join(nval(rows, rr, c).upper() for c in range(1, 41))
        if "STT" in up and ("TÊN CHI TIẾT" in up or "CCDC" in up):
            return True
    return False


def parse_reference_table(rows, start, end, max_col):
    """Generic: header row + data rows -> {headers, rows}."""
    hdr = None
    labels = []
    for r in range(start, min(end, start + 30) + 1):
        vals = [(c, nval(rows, r, c)) for c in range(1, min(max_col, 40) + 1)]
        vals = [(c, v) for c, v in vals if v]
        if len(vals) >= 3:
            up = " ".join(v.upper() for _, v in vals)
            if "STT" in up and ("TÊN CHI TIẾT" in up or "CCDC" in up):
                hdr = r
                labels = vals
                break
    if hdr is None:
        return None
    out = []
    for r in range(hdr + 1, end + 1):
        if nval(rows, r, 1) and "现场" in nval(rows, r, 1):
            break  # signature row
        row = {}
        for c, lab in labels:
            v = nval(rows, r, c)
            if v:
                row[lab] = v
        if row:
            out.append(row)
    return {"headers": [lab for _, lab in labels], "rows": out}


def parse_luutrinh(rows, max_row, max_col):
    headers = []
    for r in range(1, max_row + 1):
        if nval(rows, r, 4).lower() == "công đoạn process":
            headers.append((r, _int_or_none(val(rows, r, 1)),
                            nval(rows, r, 8)))
    chitiet = []
    blocks = []
    for i, (r, stt, ten) in enumerate(headers):
        end = headers[i + 1][0] - 1 if i + 1 < len(headers) else max_row
        fallback = not ten or ten.upper() == "THÔNG TIN"
        if fallback and _is_reference_block(rows, r, end):
            tbl = parse_reference_table(rows, r, end, max_col)
            if tbl and tbl["rows"]:
                chitiet.append(tbl)
            continue
        if fallback:
            ten = f"Công đoạn {stt}" if stt else "Công đoạn"
        b = parse_congdoan_block(rows, r, end, max_col, stt, ten)
        if fallback and not b["notes"] and not b["specs"] \
                and not b["tools"] and not b["ppe"]:
            continue  # empty template block
        blocks.append(b)
    return chitiet, blocks


# ------------------------------------------------------------ dmtg parse ----
def parse_dmtg(rows, max_row, max_col):
    hdr = None
    cols = {}
    for r in range(1, min(max_row, 30) + 1):
        vals = [nval(rows, r, c) for c in range(1, min(max_col, 15) + 1)]
        up = [v.upper() for v in vals]
        joined = " ".join(up)
        if "CÔNG ĐOẠN" in joined and ("THỜI GIAN" in joined or "T GIAN" in joined):
            hdr = r
            for c, v in enumerate(vals, start=1):
                vu = v.upper()
                if vu == "STT":
                    cols["stt"] = c
                elif "CÔNG ĐOẠN" in vu:
                    cols["ten"] = c
                elif ("T GIAN" in vu or "THỜI GIAN" in vu) and "ten" in cols \
                        and "time" not in cols and "GỌP" not in vu:
                    cols["time"] = c
                elif "NGƯỜI" in vu and "nguoi" not in cols:
                    cols["nguoi"] = c
                elif "MÁY" in vu and "may" not in cols:
                    cols["may"] = c
                elif "GHI CHÚ" in vu:
                    cols["ghichu"] = c
            break
    if hdr is None or "stt" not in cols or "ten" not in cols:
        return []
    out = []
    for r in range(hdr + 1, max_row + 1):
        stt_raw = nval(rows, r, cols["stt"])
        ten = nval(rows, r, cols["ten"])
        if not stt_raw and not ten:
            continue  # blank spacer row
        # A "TỔNG ..." row is only a subtotal separator: the same table
        # continues after it (e.g. GÒ sheets have one continuous STT 1-76
        # table split by several TỔNG subtotal rows). Skip it, don't stop.
        if "TỔNG" in ten.upper() or "TỔNG" in stt_raw.upper():
            continue
        stt = _int_or_none(stt_raw)
        if stt is None or not ten:
            continue

        def gv(k):
            return nval(rows, r, cols[k]) if k in cols else ""

        t_raw = gv("time")
        try:
            t_val = float(t_raw.replace(",", ".")) if t_raw else None
        except Exception:
            t_val = None
        out.append({
            "stt": stt, "ten": ten,
            "thoi_gian_s": t_val, "thoi_gian_raw": t_raw,
            "nguoi": gv("nguoi"), "thiet_bi": gv("may"),
            "ghi_chu": gv("ghichu"),
        })
    return out


def qtcn_stage(sheet_name):
    """Map QTCN* process-doc sheets to a stage for stage-level images."""
    s = norm(sheet_name).upper()
    if not s.startswith("QTCN"):
        return None
    if "CHẶT" in s:
        return "CHẶT"
    if "CÁN LUYỆN" in s:
        return "CÁN LUYỆN"
    if "MAY" in s:
        return "MAY"
    if "GÒ" in s:
        return "GÒ"
    if re.search(r"\bIN\b", s):
        return "IN"
    return None


# ------------------------------------------------------------------ main ---
def add_congdoan(groups, bid, stage, file_ref, sheet, stt=None, ten="",
                 notes=None, specs=None, tools=None, ppe=None,
                 thoi_gian_s=None, thoi_gian_raw="", nguoi="",
                 thiet_bi="", ghi_chu="", images=None):
    nten = re.sub(r"\s+", " ", ten.lower()).strip()
    key = (stage, nten)
    rec = groups.get(key)
    if rec is None:
        rec = groups[key] = {"brand": bid, "stage": stage, "ten": ten,
                             "stt": None, "files": [], "images": []}
    entry = {"file": file_ref, "sheet": sheet, "stt": stt}
    if notes:
        entry["dien_giai"] = notes
    if specs:
        entry["thong_so"] = specs
    if tools:
        entry["dung_cu"] = tools
    if ppe:
        entry["bao_ho"] = ppe
    if thoi_gian_s is not None or thoi_gian_raw:
        entry["thoi_gian_s"] = thoi_gian_s
        entry["thoi_gian_raw"] = thoi_gian_raw
    if nguoi:
        entry["nguoi"] = nguoi
    if thiet_bi:
        entry["thiet_bi"] = thiet_bi
    if ghi_chu:
        entry["ghi_chu"] = ghi_chu
    rec["files"].append(entry)
    if images:
        rec["images"].extend(images)


def main():
    data = {"order": STAGE_ORDER, "brands": {}}
    stats = Counter()
    for bid, bname in BRANDS:
        bdir = os.path.join(SRC, bname)
        data["brands"][bid] = {"id": bid, "name": bname, "stages": {}}
        groups = {}
        chitiet = {}
        qtcn_images = {}
        for root, _d, fs in os.walk(bdir):
            for f in sorted(fs):
                if not f.lower().endswith(".xlsx") or f.startswith("~$"):
                    continue
                path = os.path.join(root, f)
                stem = os.path.splitext(f)[0]
                file_ref = f"{bid}/{slugify(stem)}"
                try:
                    anchors = sheet_image_anchors(path)
                except Exception as e:
                    print(f"  !! anchors {f}: {e}", flush=True)
                    anchors = {}
                # stage-level images from QTCN process docs
                for sname, imgs in anchors.items():
                    qst = qtcn_stage(sname)
                    if qst:
                        qtcn_images.setdefault(qst, []).extend(
                            md5 for md5, _r, _c in imgs)
                        stats["qtcn_images"] += len(imgs)
                try:
                    wb = load_workbook(path, read_only=True, data_only=True)
                except Exception as e:
                    print(f"  !! open {f}: {e}", flush=True)
                    continue
                try:
                    for ws in wb.worksheets:
                        sname = ws.title
                        stage = stage_of(sname)
                        if not stage:
                            continue
                        max_row, max_col = ws.max_row or 0, ws.max_column or 0
                        if max_row < 5:
                            continue
                        try:
                            rows = read_all(ws, max_row, min(max_col, 71))
                        except Exception as e:
                            print(f"  !! read {f}/{sname}: {e}", flush=True)
                            continue
                        try:
                            if detect_luutrinh(rows):
                                stats["luutrinh_sheets"] += 1
                                tables, blocks = parse_luutrinh(rows, max_row, min(max_col, 71))
                                for tbl in tables:
                                    chitiet.setdefault(stage, []).append(
                                        {"file": file_ref, "sheet": sname,
                                         "headers": tbl["headers"],
                                         "rows": tbl["rows"]})
                                for b in blocks:
                                    imgs = [md5 for (md5, r0, _c0)
                                            in anchors.get(sname, [])
                                            if b["_start"] <= r0 + 1 <= b["_end"]]
                                    add_congdoan(groups, bid, stage, file_ref, sname,
                                                 stt=b["stt"], ten=b["ten"],
                                                 notes=b["notes"], specs=b["specs"],
                                                 tools=b["tools"], ppe=b["ppe"],
                                                 images=imgs)
                                    stats["luutrinh_blocks"] += 1
                            elif detect_dmtg(rows):
                                stats["dmtg_sheets"] += 1
                                for row in parse_dmtg(rows, max_row, min(max_col, 71)):
                                    add_congdoan(groups, bid, stage, file_ref, sname,
                                                 stt=row["stt"], ten=row["ten"],
                                                 thoi_gian_s=row["thoi_gian_s"],
                                                 thoi_gian_raw=row["thoi_gian_raw"],
                                                 nguoi=row["nguoi"],
                                                 thiet_bi=row["thiet_bi"],
                                                 ghi_chu=row["ghi_chu"])
                                    stats["dmtg_rows"] += 1
                        except Exception as e:
                            print(f"  !! parse {f}/{sname}: {e}", flush=True)
                finally:
                    wb.close()
                print(f"  done {file_ref}", flush=True)
        for (stage, _nten), rec in groups.items():
            st = data["brands"][bid]["stages"].setdefault(
                stage, {"stage": stage, "bang_chi_tiet": [], "cong_doan": []})
            votes = Counter(x["stt"] for x in rec["files"] if x["stt"])
            rec["stt"] = votes.most_common(1)[0][0] if votes else None
            seen = set()
            rec["images"] = [m for m in rec["images"]
                             if not (m in seen or seen.add(m))]
            st["cong_doan"].append(rec)
        for stage, lst in chitiet.items():
            st = data["brands"][bid]["stages"].setdefault(
                stage, {"stage": stage, "bang_chi_tiet": [], "cong_doan": []})
            st["bang_chi_tiet"] = lst
        for stage, imgs in qtcn_images.items():
            st = data["brands"][bid]["stages"].setdefault(
                stage, {"stage": stage, "bang_chi_tiet": [], "cong_doan": []})
            seen_q = set()
            st["qtcn_images"] = [m for m in imgs
                                 if not (m in seen_q or seen_q.add(m))]
        for st in data["brands"][bid]["stages"].values():
            st["cong_doan"].sort(key=lambda x: (x["stt"] is None, x["stt"] or 0))
    used = set()
    for bid, b in data["brands"].items():
        for st in b["stages"].values():
            for cd in st["cong_doan"]:
                base = f"{bid}-{slugify(st['stage'])}-{slugify(cd['ten'])}"
                uid, n = base, 2
                while uid in used:
                    uid = f"{base}-{n}"
                    n += 1
                used.add(uid)
                cd["id"] = uid
    json.dump(data, open(OUT, "w", encoding="utf-8"), ensure_ascii=False)
    n_cd = sum(len(st["cong_doan"]) for b in data["brands"].values()
               for st in b["stages"].values())
    print("stats:", dict(stats), flush=True)
    print(f"cong_doan groups={n_cd} -> {OUT} ({os.path.getsize(OUT)/1e6:.1f} MB)",
          flush=True)


if __name__ == "__main__":
    main()
